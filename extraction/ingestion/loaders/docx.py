"""DOCX with python-docx (contract 10 X2): a unit is a section under a heading; tables stay in document order.

The locator is the section path, for example "4 Travel > 4.2 Per diem". `page` holds the section ordinal (contract 02).
python-docx does not render automatic heading numbers, so a heading reads as it is typed.
"""
import re
from pathlib import Path

from ingestion.clean import normalise
from ingestion.layout import LIST_ITEM, SectionTracker, numbered_level
from ingestion.tables import find_unit, table_block
from ingestion.types import Block, PageResult

_HEADING_STYLE = re.compile(r"^heading\s*(\d)$", re.IGNORECASE)


def _iter_body(document):
    """Paragraphs and tables in the order they appear."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield Paragraph(child, document)
        elif tag == "tbl":
            yield Table(child, document)


def _heading_level(paragraph) -> int | None:
    style = (paragraph.style.name if paragraph.style is not None else "") or ""
    if style.lower() == "title":
        return 1
    m = _HEADING_STYLE.match(style)
    if m:
        return int(m.group(1))
    return None


def extract_docx(path: Path | str, progress=None) -> list[PageResult]:
    import docx

    document = docx.Document(str(path))
    tracker = SectionTracker()
    units: list[PageResult] = []
    current = PageResult(page=1, page_class="TEXT", method="text", section_label="")

    def start_unit():
        """A heading starts a new unit, unless the current unit holds only headings (then it continues)."""
        nonlocal current
        if all(b.type == "heading" for b in current.blocks):
            current.section_label = tracker.path
            return
        units.append(current)
        current = PageResult(page=len(units) + 1, page_class="TEXT", method="text", section_label=tracker.path)

    items: list[str] = []

    def flush_items():
        nonlocal items
        if items:
            current.blocks.append(Block("list", "\n".join(items), tracker.path))
            items = []

    for el in _iter_body(document):
        if el.__class__.__name__ == "Table":
            flush_items()
            rows = [[cell.text for cell in row.cells] for row in el.rows]
            rows = [_dedupe_merged(r) for r in rows]
            caption = ""
            if current.blocks and current.blocks[-1].type == "paragraph" and len(current.blocks[-1].text) <= 160:
                caption = current.blocks.pop().text
            unit = find_unit(caption, " ".join(rows[0]) if rows else "")
            current.blocks.append(table_block(rows, caption, unit, tracker.path))
            current.method = "text+tables"
            continue
        text = normalise(el.text)
        if not text:
            continue
        level = _heading_level(el) or (numbered_level(text) if len(text) < 80 else None)
        if level:
            flush_items()
            tracker.enter(level, text)
            start_unit()
            current.blocks.append(Block("heading", text, tracker.path, level=level))
            continue
        style = (el.style.name if el.style is not None else "").lower()
        if style.startswith("list") or LIST_ITEM.match(text):
            items.append(text)
            continue
        flush_items()
        current.blocks.append(Block("paragraph", text, tracker.path))
    flush_items()
    if current.blocks:
        units.append(current)
    for u in units:
        u.chars = sum(len(b.text) for b in u.blocks)
        u.tables = sum(1 for b in u.blocks if b.type == "table")
    if progress:
        progress("extracting", len(units), len(units))
    return units


def _dedupe_merged(cells: list[str]) -> list[str]:
    """python-docx repeats the text of a horizontally merged cell in every grid column; keep it once."""
    out = []
    for i, c in enumerate(cells):
        out.append("" if i and c == cells[i - 1] and c else c)
    return out
