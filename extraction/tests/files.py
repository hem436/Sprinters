"""Test files built in code, so the tests need no network and no binary fixtures in the repo."""
import random
from pathlib import Path

from PIL import Image, ImageDraw
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import Table, TableStyle

W, H = A4

_FONT = "Helvetica"
_FONT_BOLD = "Helvetica-Bold"
for _candidate in ("/System/Library/Fonts/Geneva.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
    if Path(_candidate).exists():          # a font with the rupee sign, so "₹ in crore" is real text in the PDF
        pdfmetrics.registerFont(TTFont("Uni", _candidate))
        _FONT = "Uni"
        break


def _write_lines(c: canvas.Canvas, lines: list[str], top: float = H - 72, size: int = 11, leading: int = 15):
    y = top
    for line in lines:
        if line.startswith("# "):
            c.setFont(_FONT_BOLD if _FONT == "Helvetica" else _FONT, 16)
            c.drawString(72, y, line[2:])
            y -= 24
            continue
        c.setFont(_FONT, size)
        if line == "":
            y -= leading                     # blank line = paragraph gap
            continue
        c.drawString(72, y, line)
        y -= leading
    return y


def text_pdf(path: Path, pages: list[list[str]], header: str | None = None, footer: bool = False) -> Path:
    """One page per list of lines. `# ` lines are drawn as large headings. Optional running header and footer."""
    c = canvas.Canvas(str(path), pagesize=A4)
    for i, lines in enumerate(pages, start=1):
        if header:
            c.setFont(_FONT, 9)
            c.drawString(72, H - 40, header)
        _write_lines(c, lines)
        if footer:
            c.setFont(_FONT, 9)
            c.drawString(W / 2, 30, f"Page {i} of {len(pages)}")
        c.showPage()
    c.save()
    return path


def noise_image(width: int = 1240, height: int = 1754, seed: int = 1) -> Image.Image:
    """A page-sized picture with scribbles: no text layer, like a scan."""
    rnd = random.Random(seed)
    img = Image.new("RGB", (width, height), "white")
    d = ImageDraw.Draw(img)
    for _ in range(300):
        x, y = rnd.randrange(width), rnd.randrange(height)
        d.line([(x, y), (x + rnd.randrange(-60, 60), y + rnd.randrange(-8, 8))], fill="black", width=2)
    return img


def mixed_pdf(path: Path) -> Path:
    """3 text pages and 1 page that is only a full-page image (contract 10 test_mixed_pdf_ocrs_only_scanned_pages)."""
    c = canvas.Canvas(str(path), pagesize=A4)
    for i in range(3):
        _write_lines(c, [f"# Section {i + 1} Travel rules", "",
                         *[f"Line {j} of the travel and expense policy text on page {i + 1}, with enough words."
                           for j in range(8)]])
        c.showPage()
    c.drawImage(ImageReader(noise_image()), 0, 0, width=W, height=H)
    c.showPage()
    c.save()
    return path


def scanned_pdf(path: Path, pages: int = 1) -> Path:
    c = canvas.Canvas(str(path), pagesize=A4)
    for i in range(pages):
        c.drawImage(ImageReader(noise_image(seed=i + 1)), 0, 0, width=W, height=H)
        c.showPage()
    c.save()
    return path


def table_pdf(path: Path) -> Path:
    """An annual-report style page: a heading, a caption, a unit line and a ruled table with printed numbers."""
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setFont(_FONT, 16)
    c.drawString(72, H - 72, "Consolidated Statement of Profit and Loss")
    c.setFont(_FONT, 11)
    c.drawString(72, H - 110, "Statement of profit and loss for the year ended 31 March 2025")
    c.drawString(72, H - 126, "(₹ in crore)")
    data = [["Particulars", "FY2024-25", "FY2023-24"],
            ["Revenue from operations", "11,256", "10,050"],
            ["Other income", "1,23,456", "98,765"],
            ["Finance costs", "(1,234)", "(1,100)"],
            ["Profit for the year", "2,345.50", "2,100.00"]]
    t = Table(data, colWidths=[200, 100, 100])
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black), ("FONTNAME", (0, 0), (-1, -1), _FONT)]))
    tw, th = t.wrapOn(c, W, H)
    t.drawOn(c, 72, H - 150 - th)
    c.setFont(_FONT, 11)
    c.drawString(72, H - 190 - th, "Revenue grew because of higher volumes in the payments segment.")
    c.showPage()
    c.save()
    return path


def big_table_pdf(path: Path, rows: int = 45) -> Path:
    """A long ruled table, for the row-split rule."""
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setFont(_FONT, 11)
    c.drawString(72, H - 60, "Segment revenue by branch (₹ in lakh)")
    data = [["Branch code", "Segment name", "FY2024-25", "FY2023-24"]]
    data += [[f"BR-{i:03d}", f"Retail payments region {i}", f"{1000 + i * 7:,}", f"{900 + i * 5:,}"]
             for i in range(rows)]
    t = Table(data, colWidths=[90, 190, 90, 90], rowHeights=15)
    t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black), ("FONTSIZE", (0, 0), (-1, -1), 8),
                           ("FONTNAME", (0, 0), (-1, -1), _FONT)]))
    tw, th = t.wrapOn(c, W, H)
    t.drawOn(c, 60, H - 80 - th)
    c.showPage()
    c.save()
    return path


def two_column_pdf(path: Path) -> Path:
    c = canvas.Canvas(str(path), pagesize=A4)
    c.setFont(_FONT, 10)
    y = H - 80
    for i in range(30):
        c.drawString(60, y, f"Left column sentence number {i:02d} here.")
        c.drawString(320, y, f"Right column sentence number {i:02d} here.")
        y -= 14
    c.showPage()
    c.save()
    return path


def encrypted_pdf(path: Path, plain: Path) -> Path:
    from pypdf import PdfReader, PdfWriter
    w = PdfWriter()
    for p in PdfReader(str(plain)).pages:
        w.add_page(p)
    w.encrypt(user_password="secret", owner_password="owner")
    with open(path, "wb") as f:
        w.write(f)
    return path


def docx_file(path: Path) -> Path:
    import docx
    d = docx.Document()
    d.add_paragraph("IT Security Standard Operating Procedure", style="Title")
    d.add_paragraph("This SOP applies to all company laptops.")
    d.add_heading("1 Scope", level=1)
    d.add_paragraph("All laptops issued by the company.")
    d.add_heading("2 Laptop controls", level=1)
    d.add_heading("2.1 Encryption", level=2)
    d.add_paragraph("Policy POL-HR-014 requires full-disk encryption on every laptop.")
    d.add_paragraph("Patches are applied within 30 days.", style="List Bullet")
    d.add_paragraph("Patch compliance by quarter (in %)")
    t = d.add_table(rows=3, cols=2)
    for r, row in enumerate([["Quarter", "Compliance"], ["Q1", "92"], ["Q2", "97"]]):
        for col, val in enumerate(row):
            t.cell(r, col).text = val
    d.save(str(path))
    return path


def png_photo(path: Path, rotate_exif: bool = False) -> Path:
    """A landscape photo; with `rotate_exif` it carries EXIF orientation 6 (stored rotated, like a phone photo)."""
    img = noise_image(800, 600)
    if rotate_exif:
        exif = img.getexif()
        exif[0x0112] = 6
        img.save(str(path.with_suffix(".jpg")), format="JPEG", exif=exif.tobytes())
        return path.with_suffix(".jpg")
    img.save(str(path), format="PNG")
    return path
